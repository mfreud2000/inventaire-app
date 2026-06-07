import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from datetime import datetime
import csv

HEADERS = ["CATEGORIE", "REFERENCE", "FOURNISSEUR", "PRODUIT", "CONDITIONNEMENT",
           "PRIX HT", "QTE M-1", "QTE M31", "INVENTAIRE VALORISE"]


def _safe_float(val):
    try:
        return float(str(val).replace(",", ".").replace(" ", ""))
    except (ValueError, TypeError):
        return 0.0


def export_to_excel(data_rows, exports_folder, month_label=None, sheet_name=None):
    folder = Path(exports_folder)
    folder.mkdir(parents=True, exist_ok=True)
    if not month_label:
        month_label = datetime.now().strftime("%Y_%m")
    filename = folder / f"Inventaire_{month_label}.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, sheet_name or f"Inventaire {month_label}", data_rows, month_label)
    wb.save(filename)
    return str(filename)


def _write_sheet(wb, ws_title, data_rows, month_label):
    """Écrit un onglet formaté dans un workbook existant."""
    ws = wb.create_sheet(title=ws_title[:31])  # Excel limite à 31 chars
    header_font  = Font(bold=True, color="FFFFFF", size=10)
    header_fill  = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alt_fill     = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    thin         = Side(style="thin", color="BBBBBB")
    cell_border  = Border(left=thin, right=thin, top=thin, bottom=thin)
    col_widths   = [16, 14, 18, 28, 16, 10, 10, 10, 18]

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = f"{ws_title} — {month_label.replace('_', '/')}"
    t.font = Font(bold=True, size=13, color="1F4E79")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for col, (h, w) in enumerate(zip(HEADERS, col_widths), start=1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = header_align; cell.border = cell_border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[2].height = 30

    total_value = 0.0
    for r_idx, row in enumerate(data_rows, start=3):
        fill   = alt_fill if r_idx % 2 == 0 else None
        padded = (row + [""] * 9)[:9]
        for col, value in enumerate(padded, start=1):
            cell = ws.cell(row=r_idx, column=col, value=value)
            if fill: cell.fill = fill
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center",
                                       horizontal="left" if col in (1,3,4) else "center")
            if col in (6, 9):
                v = _safe_float(value); cell.value = v; cell.number_format = '#,##0.00 €'
            elif col in (7, 8):
                cell.value = int(_safe_float(value))
        total_value += _safe_float(padded[8])

    total_row = len(data_rows) + 3
    ws.cell(row=total_row, column=8, value="TOTAL :").font = Font(bold=True)
    tc = ws.cell(row=total_row, column=9, value=total_value)
    tc.font = Font(bold=True, color="1F4E79"); tc.number_format = '#,##0.00 €'
    tc.fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    ws.freeze_panes = "A3"
    return total_value


def export_all_to_excel(sheets_data, exports_folder, month_label=None):
    """sheets_data = list of (sheet_name, data_rows)"""
    folder = Path(exports_folder)
    folder.mkdir(parents=True, exist_ok=True)
    if not month_label:
        month_label = datetime.now().strftime("%Y_%m")
    filename = folder / f"Inventaire_COMPLET_{month_label}.xlsx"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # supprimer la feuille vide par défaut

    grand_total = 0.0
    for sheet_name, data_rows in sheets_data:
        grand_total += _write_sheet(wb, sheet_name, data_rows, month_label)

    wb.save(filename)
    return str(filename)


def export_to_csv(data_rows, exports_folder, month_label=None):
    folder = Path(exports_folder)
    folder.mkdir(parents=True, exist_ok=True)
    if not month_label:
        month_label = datetime.now().strftime("%Y_%m")
    filename = folder / f"Inventaire_{month_label}.csv"
    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(HEADERS)
        for row in data_rows:
            writer.writerow((row + [""] * 9)[:9])
    return str(filename)
